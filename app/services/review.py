from pathlib import Path
from difflib import unified_diff
from docx import Document
from sqlalchemy import select
from sqlalchemy.orm import Session
from ..models import Submission, DocumentVersion, Finding, AuditEvent


def extract_visible_text(path: str) -> list[str]:
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return lines


def compare_versions(source: DocumentVersion, target: DocumentVersion) -> dict:
    before = extract_visible_text(source.storage_path)
    after = extract_visible_text(target.storage_path)
    changed = before != after
    diff = list(unified_diff(before, after, fromfile=f"v{source.version_number}", tofile=f"v{target.version_number}", lineterm=""))
    return {
        "source_version_id": source.id,
        "target_version_id": target.id,
        "source_version_number": source.version_number,
        "target_version_number": target.version_number,
        "text_changed": changed,
        "diff_lines": diff[:120],
        "source_sha256": source.sha256,
        "target_sha256": target.sha256,
        "change_summary": target.change_summary,
    }
