"""Controlled, formatting-only DOCX auto-fix service.

The service is deliberately allow-list based. It refuses content/semantic rules and
checks that visible document text is unchanged before a target version is accepted.
"""
from __future__ import annotations

from pathlib import Path
from copy import deepcopy
from hashlib import sha256
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

SAFE_AUTO_FIX_RULES = {
    "PAGE-001": "Set document page size to A4",
    "PAGE-002": "Set left margin to 1.5 inch",
    "PAGE-003": "Set right margin to 1 inch",
    "PAGE-004": "Set top margin to 1 inch",
    "PAGE-005": "Set bottom margin to 1 inch",
    "FONT-001": "Set run fonts to Times New Roman",
    "FONT-002": "Remove first-line paragraph indentation",
    "TAB-003": "Normalize table caption style",
    "FIG-003": "Normalize figure caption style",
    "REF-002": "Normalize reference paragraph spacing and hanging indent",
}

# Explicitly excluded even if a rule catalogue later marks them auto-fixable.
BLOCKED_RULES = {
    "TITLE-001", "TITLE-002", "TITLE-003", "TOC-001", "TOC-002",
    "XREF-001", "XREF-002", "XREF-003", "REF-001", "REF-003", "REF-004",
    "TAB-001", "TAB-002", "TAB-004", "TAB-005", "TAB-006",
    "FIG-001", "FIG-002", "FIG-004", "PRE-001", "PRE-002", "PRE-003",
}

CAPTION_RE = re.compile(r"^\s*(Table|Figure)\s+\d+(?:\.\d+)*\b", re.I)
REF_HEADING_RE = re.compile(r"^\s*(References|Bibliography)\s*$", re.I)

def _visible_text(doc: Document) -> str:
    parts: list[str] = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)

def _doc_signature(doc: Document) -> tuple[str, int, int]:
    return (_visible_text(doc), len(doc.paragraphs), len(doc.tables))

def _set_run_font(run, name="Times New Roman"):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        from docx.oxml import OxmlElement
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


def _set_caption(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    for r in p.runs:
        r.font.size = Pt(11)
        r.font.italic = True


def _normalize_references(doc):
    in_refs = False
    changed = 0
    for p in doc.paragraphs:
        if REF_HEADING_RE.match(p.text):
            in_refs = True
            continue
        if in_refs:
            if p.style and p.style.name and p.style.name.startswith("Heading"):
                break
            if p.text.strip():
                pf = p.paragraph_format
                pf.line_spacing = 1
                pf.left_indent = Inches(0.5)
                pf.first_line_indent = Inches(-0.5)
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                changed += 1
    return changed


def apply_safe_fixes(source_path: str, target_path: str, rule_ids: list[str]) -> dict:
    unique = list(dict.fromkeys(rule_ids))
    blocked = [r for r in unique if r not in SAFE_AUTO_FIX_RULES]
    if blocked:
        raise ValueError(f"Blocked or unsupported auto-fix rule(s): {', '.join(blocked)}")

    source = Path(source_path)
    target = Path(target_path)
    doc = Document(source)
    before = _doc_signature(doc)
    operations: list[str] = []

    sections = doc.sections
    if "PAGE-001" in unique:
        from docx.shared import Mm
        for sec in sections:
            sec.page_width = Mm(210)
            sec.page_height = Mm(297)
        operations.append("PAGE-001")
    if "PAGE-002" in unique:
        for sec in sections: sec.left_margin = Inches(1.5)
        operations.append("PAGE-002")
    if "PAGE-003" in unique:
        for sec in sections: sec.right_margin = Inches(1)
        operations.append("PAGE-003")
    if "PAGE-004" in unique:
        for sec in sections: sec.top_margin = Inches(1)
        operations.append("PAGE-004")
    if "PAGE-005" in unique:
        for sec in sections: sec.bottom_margin = Inches(1)
        operations.append("PAGE-005")
    if "FONT-001" in unique:
        for p in doc.paragraphs:
            for r in p.runs: _set_run_font(r)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs: _set_run_font(r)
        operations.append("FONT-001")
    if "FONT-002" in unique:
        for p in doc.paragraphs:
            p.paragraph_format.first_line_indent = Inches(0)
        operations.append("FONT-002")
    if "TAB-003" in unique or "FIG-003" in unique:
        for p in doc.paragraphs:
            m = CAPTION_RE.match(p.text)
            if not m: continue
            if m.group(1).lower() == "table" and "TAB-003" in unique:
                _set_caption(p)
            elif m.group(1).lower() == "figure" and "FIG-003" in unique:
                _set_caption(p)
        operations.append("caption-normalization")
    if "REF-002" in unique:
        _normalize_references(doc)
        operations.append("REF-002")

    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_suffix(target.suffix + ".tmp")
    doc.save(tmp)
    tmp.replace(target)

    after_doc = Document(target)
    after = _doc_signature(after_doc)
    if before != after:
        target.unlink(missing_ok=True)
        raise RuntimeError("Auto-fix safety check failed: visible document text or structure changed.")

    digest = sha256(target.read_bytes()).hexdigest()
    return {"sha256": digest, "operations": operations, "source_text_unchanged": True}
