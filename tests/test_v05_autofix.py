from pathlib import Path
from docx import Document
from app.services.autofix import apply_safe_fixes


def test_safe_autofix_preserves_visible_text(tmp_path):
    src = Path('/mnt/data/thesis_mvp_v02/Version_21_Final Complete Thesis_Disha_Wankhede.docx')
    if not src.exists():
        return
    target = tmp_path / 'v2.docx'
    before = Document(src)
    before_text = '\n'.join(p.text for p in before.paragraphs)
    result = apply_safe_fixes(str(src), str(target), ['PAGE-001', 'PAGE-002', 'PAGE-003', 'PAGE-004', 'PAGE-005', 'FONT-001', 'FONT-002', 'TAB-003', 'FIG-003', 'REF-002'])
    after = Document(target)
    after_text = '\n'.join(p.text for p in after.paragraphs)
    assert result['source_text_unchanged'] is True
    assert before_text == after_text
    assert len(before.tables) == len(after.tables)
    assert src.read_bytes() != target.read_bytes()
